#!/usr/bin/env python3
"""
DR / ASR / Runbook Excel 설계서 -> Terraform tfvars 생성기.

이 스크립트는 기존 tools/excel_to_tfvars.py가 생성하지 않는 DR 영역을 생성한다.

생성 대상:
- live/D00-foundation/resource-groups/terraform.tfvars
- live/D10-platform/dr-hub-network-jpe/terraform.tfvars
- live/D20-workload/dr-sales-prod-spoke-jpe/terraform.tfvars
- live/D30-services/dr-vm-placeholder/terraform.tfvars
- live/D40-access/dr-private-dns-records/terraform.tfvars
- live/50-asr/recovery-vault-jpe/terraform.tfvars
- live/60-operations/automation-dr/terraform.tfvars

Excel에 DR 전용 시트가 있으면 그 값을 우선 사용하고, 없으면 기본값으로 생성한다.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any, Dict, List, Optional

try:
    import openpyxl
except ImportError as exc:
    raise SystemExit("openpyxl이 필요합니다. pip install openpyxl 후 실행하십시오.") from exc


def rows_as_dicts(ws) -> List[Dict[str, Any]]:
    headers: Optional[List[str]] = None
    result: List[Dict[str, Any]] = []
    for row in ws.iter_rows(values_only=True):
        values = list(row)
        if not any(v is not None for v in values):
            continue
        if headers is None:
            headers = [str(v).strip() if v is not None else f"col_{i}" for i, v in enumerate(values)]
            continue
        item = {}
        for i in range(min(len(headers), len(values))):
            if headers[i] and values[i] not in (None, ""):
                item[headers[i]] = values[i]
        if item:
            result.append(item)
    return result


def get(row: Dict[str, Any], *keys: str, default: Any = None) -> Any:
    normalized = {str(k).strip().lower().replace(" ", "").replace("_", ""): k for k in row.keys()}
    for key in keys:
        nk = key.strip().lower().replace(" ", "").replace("_", "")
        rk = normalized.get(nk)
        if rk is not None and row[rk] not in (None, ""):
            return row[rk]
    return default


def as_bool(raw: Any, default: bool = False) -> bool:
    if raw in (None, ""):
        return default
    return str(raw).strip().lower() in {"y", "yes", "true", "1", "사용", "예", "enable", "enabled"}


def hcl_string(text: str) -> str:
    return json.dumps(str(text), ensure_ascii=False)


def hcl_key(key: Any) -> str:
    text = str(key)
    if text.replace("_", "").replace("-", "").isalnum() and not text[0].isdigit() and "-" not in text:
        return text
    return hcl_string(text)


def to_hcl(obj: Any, indent: int = 0) -> str:
    pad = " " * indent
    child = " " * (indent + 2)
    if isinstance(obj, bool):
        return "true" if obj else "false"
    if isinstance(obj, (int, float)):
        return str(obj)
    if obj is None:
        return "null"
    if isinstance(obj, str):
        return hcl_string(obj)
    if isinstance(obj, list):
        if not obj:
            return "[]"
        return "[\n" + ",\n".join(f"{child}{to_hcl(v, indent + 2)}" for v in obj) + f"\n{pad}]"
    if isinstance(obj, dict):
        if not obj:
            return "{}"
        lines = ["{"]
        for key in sorted(obj.keys()):
            lines.append(f"{child}{hcl_key(key)} = {to_hcl(obj[key], indent + 2)}")
        lines.append(f"{pad}}}")
        return "\n".join(lines)
    return hcl_string(str(obj))


def write_tfvars(path: Path, data: Dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    text = "\n\n".join(f"{key} = {to_hcl(data[key])}" for key in sorted(data.keys())) + "\n"
    path.write_text(text, encoding="utf-8")
    print(f"created: {path}")


def sheet_rows(wb, name: str) -> List[Dict[str, Any]]:
    return rows_as_dicts(wb[name]) if name in wb.sheetnames else []


def common_tags(args: argparse.Namespace, workload: str) -> Dict[str, str]:
    return {
        "project": args.project,
        "environment": "dr",
        "region_role": "dr",
        "workload": workload,
        "managed_by": "terraform",
    }


def build_d00(args: argparse.Namespace, rows: List[Dict[str, Any]]) -> Dict[str, Any]:
    if not rows:
        rows = [
            {"key": "dr_hub", "name": args.dr_hub_rg, "workload": "platform"},
            {"key": "dr_sales_prod", "name": args.dr_workload_rg, "workload": "sales"},
            {"key": "dr_asr", "name": args.dr_asr_rg, "workload": "asr"},
            {"key": "dr_automation", "name": args.dr_automation_rg, "workload": "automation"},
            {"key": "dr_inventory", "name": args.dr_inventory_rg, "workload": "inventory"},
        ]
    resource_groups = {}
    for row in rows:
        key = str(get(row, "key", "id", "구분", default=get(row, "name", default="rg"))).lower().replace("-", "_")
        name = str(get(row, "name", "ResourceGroup", "resource_group_name", default="rg-replace"))
        workload = str(get(row, "workload", "업무", default="dr"))
        resource_groups[key] = {
            "name": name,
            "location": str(get(row, "location", "Region", default=args.dr_location)),
            "tags": common_tags(args, workload),
        }
    return {"common_tags": common_tags(args, "foundation"), "resource_groups": resource_groups}


def build_d10(args: argparse.Namespace, rows: List[Dict[str, Any]]) -> Dict[str, Any]:
    subnets = {
        "firewall": {"name": "AzureFirewallSubnet", "address_prefixes": ["10.110.0.0/24"], "create_nsg": False, "associate_route_table": False},
        "dns_shared": {"name": "snet-dns-shared", "address_prefixes": ["10.110.1.0/24"], "create_nsg": True, "associate_route_table": False},
        "mgmt": {"name": "snet-mgmt", "address_prefixes": ["10.110.2.0/24"], "create_nsg": True, "associate_route_table": False},
    }
    if rows:
        subnets = {}
        for row in rows:
            key = str(get(row, "key", "id", "SubnetKey", default=get(row, "SubnetName", "name", default="subnet"))).lower().replace("-", "_")
            name = str(get(row, "SubnetName", "name", "subnet", default="snet-replace"))
            subnets[key] = {
                "name": name,
                "address_prefixes": [str(get(row, "CIDR", "address_prefix", default="10.110.0.0/24"))],
                "create_nsg": as_bool(get(row, "create_nsg", "NSG", default=True), True),
                "associate_route_table": as_bool(get(row, "associate_route_table", "RouteTable", default=False), False),
            }
    return {
        "location": args.dr_location,
        "resource_group_name": args.dr_hub_rg,
        "vnet_name": args.dr_hub_vnet_name,
        "address_space": [args.dr_hub_cidr],
        "dns_servers": [],
        "common_tags": common_tags(args, "platform"),
        "subnets": subnets,
    }


def build_d20(args: argparse.Namespace, rows: List[Dict[str, Any]]) -> Dict[str, Any]:
    subnets = {
        "web": {"name": "snet-web", "address_prefixes": ["10.120.0.0/24"], "create_nsg": True, "associate_route_table": True},
        "was": {"name": "snet-was", "address_prefixes": ["10.120.1.0/24"], "create_nsg": True, "associate_route_table": True},
        "db": {"name": "snet-db", "address_prefixes": ["10.120.2.0/24"], "create_nsg": True, "associate_route_table": True},
        "pe": {"name": "snet-pe", "address_prefixes": ["10.120.3.0/24"], "create_nsg": True, "associate_route_table": True, "private_endpoint_network_policies": "Disabled"},
        "agent": {"name": "snet-agent", "address_prefixes": ["10.120.4.0/24"], "create_nsg": True, "associate_route_table": True},
    }
    if rows:
        subnets = {}
        for row in rows:
            key = str(get(row, "key", "id", "role", default=get(row, "SubnetName", "name", default="subnet"))).lower().replace("-", "_")
            name = str(get(row, "SubnetName", "name", "subnet", default="snet-replace"))
            subnet = {
                "name": name,
                "address_prefixes": [str(get(row, "CIDR", "address_prefix", default="10.120.0.0/24"))],
                "create_nsg": as_bool(get(row, "create_nsg", "NSG", default=True), True),
                "associate_route_table": as_bool(get(row, "associate_route_table", "RouteTable", default=True), True),
            }
            if "pe" in name.lower():
                subnet["private_endpoint_network_policies"] = "Disabled"
            subnets[key] = subnet
    return {
        "location": args.dr_location,
        "resource_group_name": args.dr_workload_rg,
        "vnet_name": args.dr_workload_vnet_name,
        "address_space": [args.dr_workload_cidr],
        "dns_servers": [],
        "dr_hub_vnet_id": args.dr_hub_vnet_id,
        "dr_hub_resource_group_name": args.dr_hub_rg,
        "dr_hub_vnet_name": args.dr_hub_vnet_name,
        "dr_firewall_private_ip": None if args.dr_firewall_private_ip.lower() == "none" else args.dr_firewall_private_ip,
        "common_tags": common_tags(args, "sales"),
        "subnets": subnets,
    }


def build_d30(args: argparse.Namespace, rows: List[Dict[str, Any]]) -> Dict[str, Any]:
    if not rows:
        rows = [
            {"source_vm_name": "db01", "role": "db", "source_subnet": "snet-db", "target_subnet": "snet-db", "recovery_group": 1},
            {"source_vm_name": "was01", "role": "was", "source_subnet": "snet-was", "target_subnet": "snet-was", "recovery_group": 2},
            {"source_vm_name": "was02", "role": "was", "source_subnet": "snet-was", "target_subnet": "snet-was", "recovery_group": 2},
            {"source_vm_name": "web01", "role": "web", "source_subnet": "snet-web", "target_subnet": "snet-web", "recovery_group": 3},
            {"source_vm_name": "web02", "role": "web", "source_subnet": "snet-web", "target_subnet": "snet-web", "recovery_group": 3},
        ]
    mapping = {}
    for row in rows:
        name = str(get(row, "source_vm_name", "SourceVM", "VM Name", "VM명", default="vm"))
        mapping[name.lower().replace("-", "_")] = {
            "source_vm_name": name,
            "role": str(get(row, "role", "역할", default="app")),
            "source_subnet": str(get(row, "source_subnet", "SourceSubnet", default="snet-app")),
            "target_subnet": str(get(row, "target_subnet", "TargetSubnet", default="snet-app")),
            "recovery_group": int(get(row, "recovery_group", "RecoveryGroup", default=2)),
            "create_by": "ASR Failover",
        }
    return {"dr_vm_mapping": mapping}


def build_d40(args: argparse.Namespace, rows: List[Dict[str, Any]]) -> Dict[str, Any]:
    records = {}
    for row in rows:
        name = str(get(row, "name", "RecordName", default="")).strip()
        ip = str(get(row, "ip", "IPAddress", "records", default="")).strip()
        if not name or not ip:
            continue
        records[name] = {"name": name, "ttl": int(get(row, "ttl", default=60)), "records": [ip]}
    return {
        "dns_resource_group_name": args.dr_hub_rg,
        "private_dns_zone_name": args.private_dns_zone_name,
        "create_private_dns_zone": args.create_private_dns_zone,
        "common_tags": common_tags(args, "access"),
        "a_records": records,
    }


def build_asr_vault(args: argparse.Namespace) -> Dict[str, Any]:
    return {
        "location": args.dr_location,
        "resource_group_name": args.dr_asr_rg,
        "vault_name": args.recovery_vault_name,
        "sku": "Standard",
        "soft_delete_enabled": True,
        "common_tags": common_tags(args, "asr"),
    }


def build_automation(args: argparse.Namespace) -> Dict[str, Any]:
    return {
        "location": args.dr_location,
        "resource_group_name": args.dr_automation_rg,
        "automation_account_name": args.automation_account_name,
        "common_tags": common_tags(args, "automation"),
        "runbooks": {
            "precheck": {"name": "dr-precheck", "description": "DR pre-check", "runbook_type": "PowerShell", "script_path": "../../../runbooks/dr-precheck.ps1"},
            "health_check": {"name": "dr-health-check", "description": "DR health check", "runbook_type": "PowerShell", "script_path": "../../../runbooks/dr-health-check.ps1"},
            "dns_switch": {"name": "dr-dns-switch", "description": "DR DNS switch", "runbook_type": "PowerShell", "script_path": "../../../runbooks/dr-dns-switch.ps1"},
            "inventory_update": {"name": "dr-inventory-update", "description": "DR inventory update", "runbook_type": "PowerShell", "script_path": "../../../runbooks/dr-inventory-update.ps1"},
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--excel", required=True)
    parser.add_argument("--out", default="live")
    parser.add_argument("--project", default="azure-landing-zone-asr-runbook-ent")
    parser.add_argument("--tenant-id", required=True)
    parser.add_argument("--subscription-id", required=True)
    parser.add_argument("--dr-location", default="japaneast")
    parser.add_argument("--dr-hub-rg", default="rg-sl-dr-hub-jpe")
    parser.add_argument("--dr-workload-rg", default="rg-sl-sales-prod-dr-jpe")
    parser.add_argument("--dr-asr-rg", default="rg-sl-dr-asr-jpe")
    parser.add_argument("--dr-automation-rg", default="rg-sl-dr-automation-jpe")
    parser.add_argument("--dr-inventory-rg", default="rg-sl-dr-inventory-jpe")
    parser.add_argument("--dr-hub-vnet-name", default="vnet-sl-dr-hub-jpe")
    parser.add_argument("--dr-hub-cidr", default="10.110.0.0/16")
    parser.add_argument("--dr-workload-vnet-name", default="vnet-sl-sales-prod-dr-jpe")
    parser.add_argument("--dr-workload-cidr", default="10.120.0.0/16")
    parser.add_argument("--dr-hub-vnet-id", default="<DR_HUB_VNET_ID>")
    parser.add_argument("--dr-firewall-private-ip", default="none")
    parser.add_argument("--private-dns-zone-name", default="corp.internal")
    parser.add_argument("--create-private-dns-zone", action="store_true")
    parser.add_argument("--recovery-vault-name", default="rsv-sl-asr-jpe")
    parser.add_argument("--automation-account-name", default="aa-sl-dr-runbook-jpe")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.excel, data_only=True)
    out = Path(args.out)

    write_tfvars(out / "D00-foundation" / "resource-groups" / "terraform.tfvars", build_d00(args, sheet_rows(wb, "D00_DR_Foundation")))
    write_tfvars(out / "D10-platform" / "dr-hub-network-jpe" / "terraform.tfvars", build_d10(args, sheet_rows(wb, "D10_DR_Network")))
    write_tfvars(out / "D20-workload" / "dr-sales-prod-spoke-jpe" / "terraform.tfvars", build_d20(args, sheet_rows(wb, "D20_DR_Workload")))
    write_tfvars(out / "D30-services" / "dr-vm-placeholder" / "terraform.tfvars", build_d30(args, sheet_rows(wb, "D30_DR_VM_Mapping")))
    write_tfvars(out / "D40-access" / "dr-private-dns-records" / "terraform.tfvars", build_d40(args, sheet_rows(wb, "D40_DR_Access")))
    write_tfvars(out / "50-asr" / "recovery-vault-jpe" / "terraform.tfvars", build_asr_vault(args))
    write_tfvars(out / "60-operations" / "automation-dr" / "terraform.tfvars", build_automation(args))

    print("DR/ASR tfvars 생성 완료")


if __name__ == "__main__":
    main()
